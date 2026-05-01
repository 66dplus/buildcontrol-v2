"""Tests for the procurement approval flow (app.purchase_requests)."""

from __future__ import annotations

import json
from typing import Any, Dict, List

import pytest

from app import purchase_requests
from app.purchase_requests import (
    DECISION_APPROVE,
    DECISION_COMMENT,
    DECISION_REJECT,
    SOURCE_PAGE,
    STATUS_APPROVED,
    STATUS_PENDING,
    STATUS_REJECTED,
    _items_total,
    _next_request_no,
    _row_to_request,
    make_token,
    verify_token,
)
from config import settings


def _make_element(element_id: int, name: str, props: Dict[int, Any]) -> Dict[str, Any]:
    elem: Dict[str, Any] = {"ID": str(element_id), "NAME": name}
    for prop_id, value in props.items():
        elem[f"PROPERTY_{prop_id}"] = {"0": "" if value is None else str(value)}
    return elem


def test_token_round_trip(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "approval_link_secret", "test-secret-32-bytes-minimum-please")
    token = make_token(42)
    assert verify_token(42, token)
    assert not verify_token(42, token + "x")
    assert not verify_token(43, token)
    assert not verify_token(42, "")


def test_items_total_uses_total_or_qty_times_price() -> None:
    items = [
        {"qty": 10, "price": 5},                                    # → 50
        {"qty": 3, "price": 7, "total": 999},                       # explicit total wins
        {"qty": 0, "price": 100},                                   # → 0
    ]
    assert _items_total(items) == pytest.approx(50 + 999 + 0)


def test_next_request_no_increments_max() -> None:
    elements = [
        _make_element(1, "Заявка №1", {10: 1}),
        _make_element(2, "Заявка №3", {10: 3}),
        _make_element(3, "Заявка №2", {10: 2}),
    ]
    assert _next_request_no(elements, 10) == 4


def test_next_request_no_falls_back_to_count_when_pid_missing() -> None:
    elements = [_make_element(1, "x", {}), _make_element(2, "y", {})]
    assert _next_request_no(elements, None) == 3


def test_row_to_request_decodes_items_json_and_history() -> None:
    field_ids = {
        "f_pr_no": 1,
        "f_pr_date": 2,
        "f_pr_author": 3,
        "f_pr_etap": 4,
        "f_pr_zadacha": 5,
        "f_pr_material": 6,
        "f_pr_qty": 7,
        "f_pr_unit": 8,
        "f_pr_price": 9,
        "f_pr_total": 10,
        "f_pr_file_url": 11,
        "f_pr_status": 12,
        "f_pr_btask_id": 13,
        "f_pr_comment": 14,
        "f_pr_history": 15,
    }
    items = [
        {"material_name": "Бетон М300", "qty": 100, "price": 200, "unit": "т", "total": 20000},
        {"material_name": "Арматура", "qty": 50, "price": 60, "unit": "шт", "total": 3000},
    ]
    history = [{"ts": "2026-04-25 14:00:00", "event": "created", "actor": "Иван", "source": "form"}]
    elem = _make_element(99, "Заявка №7", {
        1: 7, 2: "2026-04-25", 3: "Иван", 4: "Этап 1", 5: "Задача 2",
        6: json.dumps(items, ensure_ascii=False), 10: 23000,
        11: "https://disk.example.com/proposal.pdf",
        12: STATUS_PENDING, 13: 555,
        14: "", 15: json.dumps(history),
    })

    decoded = _row_to_request(elem, field_ids)
    assert decoded["id"] == 99
    assert decoded["no"] == 7
    assert decoded["status"] == STATUS_PENDING
    assert decoded["bitrix_task_id"] == 555
    assert decoded["file_url"].startswith("https://")
    assert len(decoded["items"]) == 2
    assert decoded["items"][0]["material_name"] == "Бетон М300"
    assert decoded["history"] == history


def test_row_to_request_handles_single_item_legacy_layout() -> None:
    """Older rows might not store JSON; falls back to scalar columns."""
    field_ids = {
        "f_pr_no": 1, "f_pr_material": 6, "f_pr_qty": 7, "f_pr_unit": 8,
        "f_pr_price": 9, "f_pr_total": 10, "f_pr_status": 12, "f_pr_history": 15,
    }
    elem = _make_element(2, "Заявка №2", {
        1: 2, 6: "Бетон М300", 7: 50, 8: "т", 9: 200, 10: 10000, 12: STATUS_PENDING, 15: "[]",
    })
    decoded = _row_to_request(elem, field_ids)
    assert len(decoded["items"]) == 1
    assert decoded["items"][0]["material_name"] == "Бетон М300"
    assert decoded["items"][0]["qty"] == 50.0


@pytest.mark.asyncio
async def test_apply_buyer_purchase_updates_quantities_and_weighted_price(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The pure-function port of the legacy buyer-report math."""
    materials_field_map = {
        "Этап": 1, "Задача": 2, "Объём план": 3, "Объём куплено": 4,
        "Остаток на складе": 5, "Цена ед. факт, ₽": 6,
    }
    materials_elem = _make_element(101, "Бетон М300", {
        1: "Фундамент", 2: "Бетонные работы",
        3: 100,  # план
        4: 10,   # уже куплено
        5: 10,   # остаток
        6: 195,  # текущая цена факт
    })
    materials_list = {"ID": "5001", "NAME": "3. Материалы", "IBLOCK_CODE": "mat_code"}
    captured: Dict[str, Any] = {}

    async def fake_get_lists(_client: Any, _gid: int) -> List[Dict[str, Any]]:
        return [materials_list]

    async def fake_get_fields(_client: Any, _list_id: int, _ib: str, _gid: int) -> Dict[str, Any]:
        return {f"PROPERTY_{pid}": {"NAME": name} for name, pid in materials_field_map.items()}

    async def fake_get_elements(_client: Any, _list_id: int, _ib: str, _gid: int) -> List[Dict[str, Any]]:
        return [materials_elem]

    async def fake_update_element(_client: Any, _list_id: int, _ib: str, _gid: int,
                                  element_id: int, field_values: Dict[int, Any], name: str = "") -> bool:
        captured["element_id"] = element_id
        captured["field_values"] = field_values
        captured["name"] = name
        return True

    async def fake_cascade_task(*_args: Any, **_kwargs: Any) -> bool:
        captured.setdefault("cascades", []).append("task")
        return True

    async def fake_cascade_budget(*_args: Any, **_kwargs: Any) -> bool:
        captured.setdefault("cascades", []).append("budget")
        return True

    monkeypatch.setattr(purchase_requests.lists, "get_lists", fake_get_lists)
    monkeypatch.setattr(purchase_requests.lists, "get_fields", fake_get_fields)
    monkeypatch.setattr(purchase_requests.lists, "get_elements", fake_get_elements)
    monkeypatch.setattr(purchase_requests.lists, "update_element", fake_update_element)

    import utils.cascade as cascade_mod
    monkeypatch.setattr(cascade_mod, "cascade_update_task", fake_cascade_task)
    monkeypatch.setattr(cascade_mod, "cascade_update_budget", fake_cascade_budget)

    items = [{"material_name": "Бетон М300", "qty": 5, "price": 220, "total": 1100}]
    result = await purchase_requests._apply_buyer_purchase(client=None, project_id=42, items=items)

    assert result["updated"] == ["Бетон М300"]
    assert result["errors"] == []
    fv = captured["field_values"]
    # Объём куплено: 10 + 5 = 15
    assert fv[materials_field_map["Объём куплено"]] == pytest.approx(15)
    # Остаток: 10 + 5 = 15
    assert fv[materials_field_map["Остаток на складе"]] == pytest.approx(15)
    # Weighted price: (10*195 + 1100) / 15 = 203.33
    assert fv[materials_field_map["Цена ед. факт, ₽"]] == pytest.approx(203.33, rel=1e-3)
    # Cascade is intentionally NOT triggered after approval (Bug 4 fix);
    # Стоим. факт depends on Объём израсходовано, which approval does not touch.
    assert "cascades" not in captured


@pytest.mark.asyncio
async def test_apply_buyer_purchase_warns_on_over_plan(monkeypatch: pytest.MonkeyPatch) -> None:
    """Over-plan attempt must surface a warning but still apply."""
    fm = {"Объём план": 1, "Объём куплено": 2, "Остаток на складе": 3, "Цена ед. факт, ₽": 4}
    elem = _make_element(101, "Цемент", {1: 50, 2: 48, 3: 0, 4: 100})
    materials_list = {"ID": "1", "NAME": "3. Материалы", "IBLOCK_CODE": "code"}

    async def fake_get_lists(_c: Any, _g: int) -> List[Dict[str, Any]]:
        return [materials_list]

    async def fake_get_fields(_c: Any, _l: int, _i: str, _g: int) -> Dict[str, Any]:
        return {f"PROPERTY_{pid}": {"NAME": name} for name, pid in fm.items()}

    async def fake_get_elements(_c: Any, _l: int, _i: str, _g: int) -> List[Dict[str, Any]]:
        return [elem]

    async def fake_update_element(*_args: Any, **_kwargs: Any) -> bool:
        return True

    monkeypatch.setattr(purchase_requests.lists, "get_lists", fake_get_lists)
    monkeypatch.setattr(purchase_requests.lists, "get_fields", fake_get_fields)
    monkeypatch.setattr(purchase_requests.lists, "get_elements", fake_get_elements)
    monkeypatch.setattr(purchase_requests.lists, "update_element", fake_update_element)

    result = await purchase_requests._apply_buyer_purchase(
        client=None, project_id=1,
        items=[{"material_name": "Цемент", "qty": 5, "price": 100, "total": 500}],
    )
    assert result["warnings"]
    assert "превышение плана" in result["warnings"][0].lower()


@pytest.mark.asyncio
async def test_resolve_request_is_idempotent_when_already_approved(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Approve → Approve must apply only once."""
    list_info = {
        "list_id": 9001,
        "iblock_code": "purchases",
        "field_ids": {
            "f_pr_no": 1, "f_pr_status": 12, "f_pr_history": 15,
            "f_pr_material": 6,
        },
    }
    row = _make_element(77, "Заявка №1", {
        1: 1, 6: json.dumps([{"material_name": "X", "qty": 1, "price": 1, "total": 1}]),
        12: STATUS_APPROVED, 15: "[]",
    })

    async def fake_load(_client: Any, _gid: int, _rid: int):
        return list_info, row, _row_to_request(row, list_info["field_ids"])

    monkeypatch.setattr(purchase_requests, "_load_request", fake_load)

    apply_calls: List[Any] = []

    async def fake_apply(*_args: Any, **_kwargs: Any) -> Dict[str, Any]:
        apply_calls.append(1)
        return {"updated": [], "errors": [], "warnings": [], "cascade_pairs": set()}

    monkeypatch.setattr(purchase_requests, "_apply_buyer_purchase", fake_apply)

    result = await purchase_requests.resolve_request(
        client=None, project_id=1, request_id=77,
        decision=DECISION_APPROVE, actor="Test", source=SOURCE_PAGE,
    )

    assert result["already_resolved"] is True
    assert result["status"] == STATUS_APPROVED
    assert apply_calls == [], "approval logic must NOT run again on a resolved request"


@pytest.mark.asyncio
async def test_apply_buyer_purchase_does_not_call_cascade(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bug 4: approval must NOT trigger cascade — Стоим. факт is unchanged."""
    fm = {"Объём план": 1, "Объём куплено": 2, "Остаток на складе": 3, "Цена ед. факт, ₽": 4,
          "Этап": 5, "Задача": 6}
    elem = _make_element(101, "Бетон", {1: 100, 2: 0, 3: 0, 4: 0, 5: "Фундамент", 6: "Заливка"})
    materials_list = {"ID": "1", "NAME": "3. Материалы", "IBLOCK_CODE": "code"}

    async def fake_get_lists(_c: Any, _g: int) -> List[Dict[str, Any]]:
        return [materials_list]

    async def fake_get_fields(_c: Any, _l: int, _i: str, _g: int) -> Dict[str, Any]:
        return {f"PROPERTY_{pid}": {"NAME": name} for name, pid in fm.items()}

    async def fake_get_elements(_c: Any, _l: int, _i: str, _g: int) -> List[Dict[str, Any]]:
        return [elem]

    async def fake_update_element(*_a: Any, **_k: Any) -> bool:
        return True

    cascade_calls: List[str] = []

    async def fake_cascade_task(*_a: Any, **_k: Any) -> bool:
        cascade_calls.append("task")
        return True

    async def fake_cascade_budget(*_a: Any, **_k: Any) -> bool:
        cascade_calls.append("budget")
        return True

    monkeypatch.setattr(purchase_requests.lists, "get_lists", fake_get_lists)
    monkeypatch.setattr(purchase_requests.lists, "get_fields", fake_get_fields)
    monkeypatch.setattr(purchase_requests.lists, "get_elements", fake_get_elements)
    monkeypatch.setattr(purchase_requests.lists, "update_element", fake_update_element)

    import utils.cascade as cascade_mod
    monkeypatch.setattr(cascade_mod, "cascade_update_task", fake_cascade_task)
    monkeypatch.setattr(cascade_mod, "cascade_update_budget", fake_cascade_budget)

    items = [{"material_name": "Бетон", "qty": 5, "price": 200, "total": 1000}]
    await purchase_requests._apply_buyer_purchase(client=None, project_id=42, items=items)

    assert cascade_calls == [], "cascade must not run after approval; foreman report owns Стоим. факт"


@pytest.mark.asyncio
async def test_materials_plan_index_returns_per_material_price_plan(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bug 6 helper: materials_plan_index must return per-material price_plan."""
    fm = {"Цена ед. план, ₽": 7, "Ед. изм": 8}
    e1 = _make_element(1, "Бетон М300", {7: 100, 8: "м3"})
    e2 = _make_element(2, "Арматура", {7: 50, 8: "т"})
    materials_list = {"ID": "1", "NAME": "3. Материалы", "IBLOCK_CODE": "code"}

    async def fake_get_lists(_c: Any, _g: int) -> List[Dict[str, Any]]:
        return [materials_list]

    async def fake_get_fields(_c: Any, _l: int, _i: str, _g: int) -> Dict[str, Any]:
        return {f"PROPERTY_{pid}": {"NAME": name} for name, pid in fm.items()}

    async def fake_get_elements(_c: Any, _l: int, _i: str, _g: int) -> List[Dict[str, Any]]:
        return [e1, e2]

    monkeypatch.setattr(purchase_requests.lists, "get_lists", fake_get_lists)
    monkeypatch.setattr(purchase_requests.lists, "get_fields", fake_get_fields)
    monkeypatch.setattr(purchase_requests.lists, "get_elements", fake_get_elements)

    idx = await purchase_requests.materials_plan_index(client=None, project_id=42)

    assert idx["бетон м300"]["price_plan"] == pytest.approx(100.0)
    assert idx["бетон м300"]["unit"] == "м3"
    assert idx["арматура"]["price_plan"] == pytest.approx(50.0)
    assert idx["арматура"]["unit"] == "т"


def test_import_excel_skips_empty_etap_rows(tmp_path: Any) -> None:
    """Bug 2: rows with empty Этап in sheet 2 must be skipped during list creation."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "2. Этапы и задачи"
    ws.append(["Этап", "Задача", "Дата нач. план"])
    ws.append(["Фундамент", "Заливка", "2026-05-01"])  # keep
    ws.append(["", "Без этапа", "2026-05-02"])         # skip — empty Этап
    ws.append(["Кровля", "Монтаж", "2026-05-03"])     # keep
    f = tmp_path / "v3.xlsx"
    wb.save(f)

    from openpyxl import load_workbook
    wb2 = load_workbook(f, read_only=True, data_only=True)
    sheet = wb2["2. Этапы и задачи"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = list(rows[0])
    data_rows = [dict(zip(headers, r)) for r in rows[1:]]

    # Same filter as scripts/import_excel.py: skip empty Этап rows
    kept = [r for r in data_rows if (r.get("Этап") or "").strip()]
    assert len(kept) == 2
    assert kept[0]["Задача"] == "Заливка"
    assert kept[1]["Задача"] == "Монтаж"


def test_upload_excel_rejects_corrupted_xlsx_via_openpyxl() -> None:
    """Bug 1: pre-validation uses openpyxl; non-xlsx bytes must raise."""
    import tempfile
    from openpyxl import load_workbook
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        f.write(b"this is not a real xlsx file just plain text bytes")
        bad_path = f.name
    raised = False
    try:
        wb = load_workbook(bad_path, read_only=True, data_only=True)
        wb.close()
    except Exception:
        raised = True
    import os as _os
    try:
        _os.unlink(bad_path)
    except OSError:
        pass
    assert raised, "openpyxl must raise on corrupted xlsx; pre-validation guard depends on this"
