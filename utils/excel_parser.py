"""
Excel file parser for BuildControl import templates.
"""

from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook


def parse_excel_file(file_path: str) -> Dict[str, List[Dict[str, Any]]]:
    """
    Parse Excel file and extract all 4 sheets as structured data.

    Template structure:
    - Row 1: Title row
    - Row 2-3: Headers row (find row with actual column names)
    - Row 4+: Data rows

    Args:
        file_path: Path to Excel file

    Returns:
        Dict mapping sheet names to list of row dicts
        Example: {
            "1. Бюджет": [{"Статья затрат": "СМР", "Кол-во план": 100, ...}, ...],
            "2. Материалы": [...],
            ...
        }

    Raises:
        FileNotFoundError: If file doesn't exist
        ValueError: If workbook structure is invalid
    """
    file_path_obj = Path(file_path)
    if not file_path_obj.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    workbook = load_workbook(file_path, data_only=True)
    result: Dict[str, List[Dict[str, Any]]] = {}

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]

        # Find header row (first row with multiple non-None values)
        header_row_idx = 1
        headers: List[str] = []

        for row_idx in range(1, 5):  # Check first 4 rows for headers
            row_values = list(worksheet.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
            non_none_count = sum(1 for v in row_values if v is not None)

            # If row has 2+ non-None values, it's likely the header row
            if non_none_count >= 2:
                headers = [str(v).strip() if v else "" for v in row_values[:non_none_count]]
                headers = [h for h in headers if h]  # Remove empty strings
                header_row_idx = row_idx
                break

        if not headers:
            raise ValueError(f"Sheet '{sheet_name}' has no valid headers")

        # Parse data rows starting after header row
        rows: List[Dict[str, Any]] = []
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
            row_data: Dict[str, Any] = {}
            for col_idx, value in enumerate(row):
                if col_idx < len(headers):
                    header = headers[col_idx]
                    if header:  # Only add if header is not empty
                        row_data[header] = value

            # Skip completely empty rows
            if any(v is not None for v in row_data.values()):
                rows.append(row_data)

        result[sheet_name] = rows

    workbook.close()
    return result


def get_project_name(file_path: str) -> str:
    """
    Extract project name from Excel filename (without extension).

    Args:
        file_path: Path to Excel file

    Returns:
        Project name (filename without extension)
    """
    return Path(file_path).stem


def get_sheet_structure(file_path: str) -> Dict[str, Tuple[int, List[str]]]:
    """
    Get structure info (row count and headers) for each sheet.

    Args:
        file_path: Path to Excel file

    Returns:
        Dict mapping sheet name to (row_count, headers_list)
    """
    data = parse_excel_file(file_path)
    result: Dict[str, Tuple[int, List[str]]] = {}

    for sheet_name, rows in data.items():
        headers = list(rows[0].keys()) if rows else []
        result[sheet_name] = (len(rows), headers)

    return result
